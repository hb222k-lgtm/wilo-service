"""문서 내보내기 — 엑셀(.xlsx) / 한글(.hwpx)

문서마다 레이아웃을 한 번만 정의하고(블록 목록), 같은 레이아웃을 엑셀과 한글로 각각 그린다.
  블록: ('p', 문단) · ('s', 여백mm) · ('t', 표) · ('br',) 페이지 나눔
  표:   widths=[열 너비 mm...], rows=[(행 높이 mm, [셀...])]  — 셀은 HTML처럼 병합된 칸은 생략
"""
import io, os, zipfile, random
from xml.sax.saxutils import escape
from PIL import Image, ImageOps

PAGE_W = 180  # A4 본문 폭(mm), 좌우 여백 15mm
STAMP = 15    # 직인 크기(mm)

TYPE_NAME = {'quote': '견적서', 'statement': '거래명세서', 'contract': '공사도급표준계약서', 'completion': '작업완료확인서'}
GRAY = '#EFEFEF'
BLUE = '#DFE5EE'


# ── 공통 ──────────────────────────────────────────────────────────────────────

def won(n):
    return f'{int(n or 0):,}'


def kor_num(n):
    """숫자 → 한글 금액 (2420000 → 이백사십이만)"""
    n = int(n or 0)
    if not n:
        return '영'
    D, S, B = '_일이삼사오육칠팔구', ['', '십', '백', '천'], ['', '만', '억', '조']
    out, g = '', 0
    while n > 0:
        part = n % 10000
        if part:
            s = ''
            for i, c in enumerate(f'{part:04d}'):
                v, u = int(c), 3 - i
                if v:
                    s += ('' if v == 1 and u > 0 else D[v]) + S[u]
            if part == 1 and g > 0:
                s = '일'
            out = s + B[g] + out
        n //= 10000
        g += 1
    return out


def item_tax(amount, vat_mode):
    """부가세: 'excluded' 별도(10%) · 'included' 포함 · 'none' 없음(현금 등)"""
    return 0 if vat_mode == 'none' else round(amount * 0.1)


def vat_label(vat_mode):
    return {'included': '부가세포함', 'none': '부가세 없음'}.get(vat_mode, '부가세별도')


def sum_items(items, vat_mode=None):
    supply = tax = 0
    for it in items:
        a = round(float(it.get('qty') or 0) * float(it.get('price') or 0))
        supply += a
        tax += item_tax(a, vat_mode)
    return supply, tax


def doc_bank(doc, co):
    """문서에 고른 입금계좌 (명세서는 안 고르면 설정의 첫 번째 계좌)"""
    d = doc['data']
    if 'bank' in d:
        return d['bank'] or ''
    banks = [b.strip() for b in (co.get('bank') or '').split('\n') if b.strip()]
    return banks[0] if doc['doc_type'] == 'statement' and banks else ''


def ymd(d, blank=False):
    if not d:
        return '        년      월      일' if blank else ''
    y, m, dd = d.split('-')
    return f'{y}년 {m}월 {dd}일'


def ymd_short(d):
    if not d:
        return ''
    y, m, dd = d.split('-')
    return f'{y}년 {int(m)}월 {int(dd)}일'


def spaced(s):
    return ' '.join(s or '')


def vertical(s):
    return '\n'.join(c for c in s if c != ' ')


def C(text='', **kw):
    kw['text'] = '' if text is None else str(text)
    return kw


def P(text, **kw):
    kw['text'] = text
    return ('p', kw)


def T(widths, rows, border=True):
    return ('t', {'widths': widths, 'rows': rows, 'border': border})


# ── 문서별 레이아웃 ───────────────────────────────────────────────────────────

def lab(text, **kw):
    return C(text, bg=GRAY, bold=True, align='CENTER', **kw)


def layout_quote(doc, co):
    d, c = doc['data'], doc['data'].get('client') or {}
    items = [it for it in d.get('items') or [] if it.get('name') or it.get('price')]
    supply, tax = sum_items(items, d.get('vat_mode'))
    total = supply + tax
    tel = ('TEL ' + c['tel'] if c.get('tel') else '') + (' / FAX ' + c['fax'] if c.get('fax') else '')
    co_tel = f"TEL {co['tel']}" + (f" / FAX {co['fax']}" if co.get('fax') else '')
    none = (0, 0, 0, 0)
    rows = [(7.5, [lab('수 신'), C(doc['client_name'], bold=True), C(b=none), lab('상 호'), C(co['name'], b=(1, 0, 1, 1)),
                   C('(인)', align='CENTER', b=(0, 1, 1, 1), stamp=True)]),
            (7.5, [lab('연락처'), C(tel), C(b=none), lab('주 소'), C(co['address'], cs=2)]),
            (7.5, [lab('공사명'), C(doc.get('title')), C(b=none), lab('사업자번호'), C(co['biz_no'], cs=2)]),
            (7.5, [lab('유효기간'), C(d.get('valid_until')), C(b=none), lab('연락처', rs=2), C(f"{co_tel}\nEMAIL  {co['email']}", rs=2, cs=2)]),
            (7.5, [C('1. 귀사의 일익 번창하심을 기원합니다.\n2. 하기와 같이 견적드리오니 검토하기 바랍니다.', cs=2, rs=2), C(b=none)]),
            (7.5, [C(b=none), lab('담당자'), C(f"{co['ceo']}   {co.get('mobile', '')}", cs=2)])]
    item_rows = [(7.6, [lab(h) for h in ['품목명', '규격', '수량', '단가', '공급가액', '부가세']])]
    for i in range(max(14, len(items))):
        it = items[i] if i < len(items) else None
        if not it:
            item_rows.append((7, [C() for _ in range(6)]))
            continue
        a = round(float(it.get('qty') or 0) * float(it.get('price') or 0))
        item_rows.append((7, [C(it.get('name')), C(it.get('spec'), align='CENTER'), C(it.get('qty'), align='CENTER'),
                              C(num=it.get('price') or 0, align='RIGHT'), C(num=a, align='RIGHT'), (C('-', align='RIGHT') if d.get('vat_mode') == 'none' else C(num=item_tax(a, d.get('vat_mode')), align='RIGHT'))]))
    return [
        P('견 적 서', size=22, bold=True, underline=True, align='CENTER', title=True),
        ('s', 3),
        T([90, 90], [(8, [C(f"일 자 : {(doc['doc_date'] or '').replace('-', '/')}", valign='BOTTOM'),
                          C(img=('logo', 25, 7), align='RIGHT', valign='BOTTOM')])], border=False),
        T([20, 68, 4, 22, 54, 12], rows),
        ('s', 3),
        T([120, 60], [(9, [C(f'금 액 : {kor_num(total)}원 정', bold=True, size=11.5, b=(1, 0, 1, 1)),
                           C(f'(₩ {won(total)}원)', bold=True, size=11.5, align='RIGHT', b=(0, 1, 1, 1))])]),
        ('s', 3),
        T([69, 16, 15, 26, 30, 24], item_rows),
        ('s', 2),
        T([22, 38, 14, 40, 16, 50], [(9, [lab('공급가액'), C(num=supply, align='RIGHT'), lab('VAT'), C(num=tax, align='RIGHT'),
                                          lab('합계'), C(num=total, align='RIGHT')])]),
        ('s', 2),
        T([15, 165], [(22, [lab('비고'), C(d.get('note'), valign='BOTTOM')])]),
    ] + bank_block(doc, co)


def layout_statement(doc, co):
    d, c = doc['data'], doc['data'].get('client') or {}
    items = [it for it in d.get('items') or [] if it.get('name') or it.get('price')]
    supply, tax = sum_items(items, d.get('vat_mode'))
    total = supply + tax
    none = (0, 0, 0, 0)
    vlab = lambda s: C(vertical(s), rs=5, bg=GRAY, bold=True, align='CENTER', size=9.5)
    rows = [(7.5, [vlab('공급받는자'), lab('상호'), C(doc['client_name'], cs=3, bold=True), C(b=none),
                   vlab('공급자'), lab('상호'), C(co['name'], cs=3)]),
            (7.5, [lab('사업자번호', size=9), C(c.get('biz_no'), size=9), lab('성명', size=9), C(c.get('ceo'), size=9), C(b=none),
                   lab('사업자번호', size=9), C(co['biz_no'], size=9), lab('성명', size=9), C(co['ceo'], size=9, stamp=True)]),
            (7.5, [lab('주소'), C(c.get('address'), cs=3, size=9), C(b=none), lab('주소'), C(co['address'], cs=3, size=9)]),
            (7.5, [lab('연락처'), C(c.get('tel'), cs=3), C(b=none), lab('연락처'), C(co['tel'], cs=3)]),
            (7.5, [lab('공사명'), C(doc.get('title'), cs=3, size=9), C(b=none), lab('EMAIL'), C(co['email'], cs=3)])]
    item_rows = [(7.6, [lab(h) for h in ['순번', '품목명', '규격', '수량', '단가', '공급가액', '부가세']])]
    for i in range(max(14, len(items))):
        it = items[i] if i < len(items) else None
        if not it:
            item_rows.append((7, [C() for _ in range(7)]))
            continue
        a = round(float(it.get('qty') or 0) * float(it.get('price') or 0))
        item_rows.append((7, [C(i + 1, align='CENTER'), C(it.get('name')), C(it.get('spec'), align='CENTER'), C(it.get('qty'), align='CENTER'),
                              C(num=it.get('price') or 0, align='RIGHT'), C(num=a, align='RIGHT'), (C('-', align='RIGHT') if d.get('vat_mode') == 'none' else C(num=item_tax(a, d.get('vat_mode')), align='RIGHT'))]))
    blocks = [
        P('거 래 명 세 서', size=22, bold=True, underline=True, align='CENTER', title=True),
        ('s', 3),
        T([20, 50, 110], [(7.5, [lab('일 자'), C((doc['doc_date'] or '').replace('-', '/')),
                                 C(img=('logo', 25, 7), align='RIGHT', valign='BOTTOM', b=none)])]),
        T([7, 19, 30, 10, 21, 3, 7, 19, 30, 10, 24], rows),
        ('s', 3),
        T([120, 60], [(6, [C('  다음과 같이 계산합니다.', cs=2, size=9.5, b=(1, 1, 1, 0))]),
                      (8, [C(f'금  액 : {kor_num(total)}원 정', bold=True, size=11.5, b=(1, 0, 0, 1)),
                           C(f'(₩{won(total)})', bold=True, size=11.5, align='RIGHT', b=(0, 1, 0, 1))])]),
        T([11, 62, 15, 13, 25, 29, 25], item_rows),
        ('s', 2),
        T([20, 30, 12, 28, 13, 30, 13, 34], [(9, [lab('공급가액'), C(num=supply, align='RIGHT'), lab('VAT'), C(num=tax, align='RIGHT'),
                                                  lab('합계'), C(num=total, align='RIGHT'), lab('확인'), C('인', align='RIGHT')])]),
        ('s', 2),
        T([20, 160], [(15, [lab('비 고'), C(d.get('note'))])]),
    ]
    return blocks + bank_block(doc, co)


def bank_block(doc, co):
    bank = doc_bank(doc, co)
    return [('s', 2), T([20, 160], [(9, [lab('입금계좌'), C(bank, bold=True)])])] if bank else []


def layout_contract(doc, co):
    d, c = doc['data'], doc['data'].get('client') or {}
    amt = int(doc.get('supply_amount') or 0)
    vat = vat_label(d.get('vat_mode'))
    ctr = lambda s, **kw: C(s, align='CENTER', **kw)
    lbl = lambda s, **kw: C(s, align='CENTER', role='label', **kw)  # 항목 이름 칸 (새 배치에서 색칠)
    # 서명란: 왼쪽 "갑"(3칸) · 오른쪽 "을"(3칸, 마지막 칸은 (인) 자리)
    L, R = (1, 1, 0, 0), (1, 1, 0, 0)
    sign = [(9, [C('"갑"', cs=3, b=(1, 1, 1, 0), role='who'), C('"을"', cs=3, b=(1, 1, 1, 0), role='who')]),
            (8.5, [C(doc['client_name'], cs=3, b=L), C(f'주    소 : {co["address"]}', cs=3, b=R, size=9.5)]),
            (8.5, [C(f'주    소 : {c.get("address", "")}', cs=3, b=L), C(f'상    호 : {co["name"]}', cs=3, b=R)]),
            (8.5, [C(f'대 표 자 : {spaced(c.get("ceo"))} ( 인 )', cs=3, b=L), C(f'대 표 자 : {spaced(co["ceo"])}', cs=2, b=(1, 0, 0, 0)),
                   C('(인)', align='CENTER', b=(0, 1, 0, 0), stamp=True)]),
            (8.5, [C(f'전    화 : {c.get("tel", "")}', cs=3, b=L), C(f'전    화 : {co["tel"]}', cs=3, b=R)]),
            (5, [C(cs=3, b=(1, 1, 0, 1)), C(cs=3, b=(1, 1, 0, 1))])]
    rows = [
        (9.5, [C('공사도급표준계약서', cs=3, rs=2, size=17, bold=True, align='CENTER', title=True), lbl('계약일자', cs=3)]),
        (9.5, [ctr(ymd(doc['doc_date']), cs=3)]),
        (9.5, [ctr('계\n약\n자', rs=4, bold=True, role='sec'), lbl('발주처'), C(f"○ {doc['client_name']}", cs=4)]),
        (9.5, [lbl('계약대상자', rs=3), C(f"○ 상호 또는 법인 명칭 : {co['name']}"), C(f"○ 사업자등록번호 : {co['biz_no']}", cs=3, size=9.5)]),
        (9.5, [C(f"○ 주소 : {co['address']}"), C(f"○ 전화번호 : {co['tel']}", cs=3)]),
        (9.5, [C(f"○ 대표자 : {spaced(co['ceo'])}"), C(cs=3)]),
        (9.5, [ctr('계\n약\n내\n용', rs=8, bold=True, role='sec'), lbl('공 사 명'), C(doc.get('title'), cs=4)]),
        (9.5, [lbl('계 약 금'), C(f'일금 {kor_num(amt)} 원정 ( ₩{won(amt)}원/{vat} )', cs=4)]),
        (9.5, [lbl('계약이행보증금'), C(d.get('perf_bond'), cs=4)]),
        (9.5, [lbl('하자이행보증금'), C(d.get('defect_bond'), cs=4)]),
        (9.5, [lbl('하자담보기간'), C(d.get('warranty'), cs=4)]),
        (9.5, [lbl('지체상금율'), C(d.get('delay_rate'), cs=4)]),
        (9.5, [lbl('착공년월일'), ctr(ymd(d.get('start_date'), True)), lbl('준공년월일'), ctr(ymd(d.get('end_date'), True), cs=2)]),
        (9.5, [lbl('기타사항'), C(d.get('etc_note'), cs=4)]),
        (22, [C(role='text', text='도급인("갑")과 수급인("을")은 상기와 같이 공사계약을 체결하고 신의에 성실히 계약상의 의무를 이행할 것을 확약하며, '
                '이 계약의 증서로서 계약서를 작성하여 당사자 간 기명, 날인 후 각각 1통씩 보관한다.', cs=6)]),
        (9.5, [C(f"붙임서류 : {d.get('attachments') or ''}", cs=3), C(cs=3)]),
        (12, [C(ymd(doc['doc_date']) + '   ', cs=6, align='RIGHT')]),
    ] + sign
    return [T([13, 30, 73, 27, 25, 12], rows)]


def layout_completion(doc, co):
    d = doc['data']
    amt = int(doc.get('supply_amount') or 0)
    vat = vat_label(d.get('vat_mode'))
    none = (0, 0, 0, 0)
    info = lambda k, v: (11, [C(k, bg=BLUE, b=none, size=11), C(v, bg=BLUE, b=none, size=11)])
    blocks = [
        ('s', 22),
        P('작업완료확인서', size=22, align='CENTER', title=True),
        ('s', 10),
        T([30, 150], [info('○ 공   사  명 :', doc.get('title')),
                      info('○ 계약금액 :', f'일 금 {kor_num(amt)} 원 정 (₩{won(amt)}), {vat}' if amt else ''),
                      info('○ 계약일자 :', ymd_short(d.get('contract_date'))),
                      info('○ 착공일자 :', ymd_short(d.get('start_date'))),
                      info('○ 준공일자 :', ymd_short(d.get('end_date')))], border=False),
        ('s', 30),
        P('    ❖    별첨 : 사진자료', size=12),
        ('s', 12),
        P((doc['doc_date'] or '').replace('-', '.'), size=13, align='CENTER'),
        ('s', 32),
        T([96, 32, 12, 40], [(8, [C(), C(f"주  소 : {co['address']}", cs=3, size=10.5)]),
                             (8, [C(), C(f"회사명 : {co['name']}", cs=3, size=10.5)]),
                             (8, [C(), C(f"성  명 : {co['ceo']}", size=10.5), C('(인)', align='CENTER', size=10.5, stamp=True), C()])],
          border=False),
    ]
    if (d.get('photo_layout') or 'rows') == 'grid':
        return blocks + completion_grid(d.get('photos') or [])
    groups = [g for g in d.get('groups') or [] if g.get('photos')]
    for p in range(0, len(groups), 3):
        rows = []
        for g in groups[p:p + 3]:
            photos = (g['photos'] + [None, None, None])[:3]
            rows.append((70, [C(vertical(g.get('title') or ''), rs=2, bold=True, align='CENTER', b=none, size=10.5)] +
                         [C(img=('photo', ph['folder'], ph['filename'], 54, 68), align='CENTER', b=none) if ph else C(b=none) for ph in photos]))
            rows.append((7, [C(ph.get('caption') if ph else '', align='CENTER', b=none) for ph in photos]))
            rows.append((4, [C(b=none, cs=4)]))
        blocks += [('br',), P('작업사진', size=14, bold=True, align='CENTER'), ('s', 5), T([12, 56, 56, 56], rows, border=False)]
    return blocks


def grid_shape(n):
    """사진 수에 따라 (열 수, 사진 폭mm, 사진 높이mm) — 한 페이지 최대 12장"""
    cols = 2 if n <= 4 else 3 if n <= 9 else 4
    rows = -(-n // cols)
    w = 178 / cols - 4
    h = min(w * 1.3, 240 / rows - 9)
    return cols, w, h


def completion_grid(photos):
    """작업사진: 사진만 격자로 모아서 (항목 제목 · 설명 없음)"""
    none = (0, 0, 0, 0)
    blocks = []
    for p in range(0, len(photos), 12):
        page = photos[p:p + 12]
        cols, w, h = grid_shape(len(page))
        rows = []
        for r in range(0, len(page), cols):
            line = page[r:r + cols] + [None] * (cols - len(page[r:r + cols]))
            rows.append((h + 2, [C(img=('photo', ph['folder'], ph['filename'], w, h), align='CENTER', b=none) if ph else C(b=none)
                                 for ph in line]))
            rows.append((4, [C(b=none, cs=cols)]))
        blocks += [('br',), P('작업사진', size=14, bold=True, align='CENTER'), ('s', 5),
                   T([178 / cols] * cols, rows, border=False)]
    return blocks


# ── 새 배치 (레터형 · 카드형 · 사이드바형) — print-layouts.js와 같은 구성 ────────
# 색은 자리표시(@accent 등)로 두고 apply_theme에서 시안 색으로 바꾼다.
ACCENT, SOFT, LABBG, LABFG = '@accent', '@soft', '@lab_bg', '@lab_fg'
WHITE, MUTED = '#FFFFFF', '#666666'
NONE, HAIR = (0, 0, 0, 0), (0, 0, 0, 1)

TRADE = {
    'quote': {'title': '견 적 서', 'plain': '견적서', 'en': 'QUOTATION', 'numbered': False,
              'greet': '귀사의 일익 번창하심을 기원합니다. 아래와 같이 견적드리오니 검토 바랍니다.'},
    'statement': {'title': '거 래 명 세 서', 'plain': '거래명세서', 'en': 'STATEMENT', 'numbered': True,
                  'greet': '다음과 같이 계산합니다.'},
}


def trade_meta(kind, doc, c, d):
    if kind == 'quote':
        tel = ' · '.join(x for x in [c.get('tel') and 'TEL ' + c['tel'], c.get('fax') and 'FAX ' + c['fax']] if x)
        rows = [('공사명', doc.get('title')), ('유효기간', d.get('valid_until')), ('연락처', tel)]
    else:
        rows = [('공사명', doc.get('title')), ('사업자번호', c.get('biz_no')), ('대표자', c.get('ceo')),
                ('주소', c.get('address')), ('연락처', c.get('tel'))]
    return [(k, v) for k, v in rows if v]


def vat_word(d):
    return '부가세 없음' if d.get('vat_mode') == 'none' else 'VAT 포함'


def doc_meta(doc, en=''):
    parts = [en, doc.get('doc_no'), (doc.get('doc_date') or '').replace('-', '.')]
    return ' · '.join(p for p in parts if p)


def trade_items(doc):
    return [it for it in doc['data'].get('items') or [] if it.get('name') or it.get('price')]


def item_widths(numbered, compact=False):
    if compact:  # 사이드바형 본문 폭 128mm
        return [8, 40, 11, 10, 18, 22, 19] if numbered else [48, 11, 10, 18, 22, 19]
    return [9, 62, 14, 13, 24, 30, 28] if numbered else [71, 14, 13, 24, 30, 28]


def item_rows(items, d, numbered, style, min_rows=10):
    """품목표 행 — style 'letter'(머리칸 색 · 줄무늬) · 'card'(강조색 글씨 · 가는 선) · 'side'"""
    head = (['No'] if numbered else []) + ['품목명', '규격', '수량', '단가', '공급가액', '부가세']
    al = (['CENTER'] if numbered else []) + ['LEFT', 'CENTER', 'CENTER', 'RIGHT', 'RIGHT', 'RIGHT']
    if style == 'card':
        rows = [(7.5, [C(h, bold=True, size=8.5, color=ACCENT, b=HAIR, align='CENTER' if a == 'LEFT' else a) for h, a in zip(head, al)])]
    else:
        rows = [(7.5, [C(h, bold=True, size=9, bg=LABBG, color=LABFG, b=HAIR, align='CENTER' if a == 'LEFT' else a) for h, a in zip(head, al)])]
    none_vat = d.get('vat_mode') == 'none'
    for i in range(max(min_rows, len(items))):
        bg = SOFT if style == 'letter' and i % 2 == 1 else None
        if i >= len(items):
            rows.append((7, [C(b=HAIR, bg=bg) for _ in head]))
            continue
        it = items[i]
        a = round(float(it.get('qty') or 0) * float(it.get('price') or 0))
        vals = ([C(i + 1)] if numbered else []) + [C(it.get('name')), C(it.get('spec')), C(it.get('qty')),
                C(num=it.get('price') or 0), C(num=a), C('-') if none_vat else C(num=item_tax(a, d.get('vat_mode')))]
        for cell, a_ in zip(vals, al):
            cell.update(align=a_, b=HAIR, size=9.5, **({'bg': bg} if bg else {}))
        rows.append((7, vals))
    return rows


def sum_rows(supply, tax, ncols, size=9.5, grand=11.5):
    """합계 3줄 — 마지막 두 칸에 이름 · 금액"""
    pad = [C(cs=ncols - 2, b=NONE)] if ncols > 2 else []
    return [(7, pad + [C('공급가액', b=HAIR, size=size), C(num=supply, align='RIGHT', b=HAIR, size=size)]),
            (7, [C(cs=ncols - 2, b=NONE)] + [C('부가세', b=HAIR, size=size), C(num=tax, align='RIGHT', b=HAIR, size=size)]),
            (8.5, [C(cs=ncols - 2, b=NONE)] + [C('합계', bold=True, color=ACCENT, b=HAIR, size=11),
                                               C(f'₩ {won(supply + tax)}', bold=True, color=ACCENT, align='RIGHT', b=HAIR, size=grand)])]


def letter_head(co, title, doc):
    info2 = ' · '.join(x for x in [f"TEL {co.get('tel', '')}", co.get('fax') and f"FAX {co['fax']}", co.get('email')] if x)
    return [T([110, 70], [
        (10, [C(co['name'], size=18, bold=True, color=ACCENT, valign='BOTTOM'), C(img=('logo', 25, 7), align='RIGHT', valign='BOTTOM')]),
        (5, [C(co.get('address'), size=8.5, color=MUTED), C(title, rs=2, size=20, bold=True, color=ACCENT, align='RIGHT', valign='BOTTOM')]),
        (5, [C(info2, size=8.5, color=MUTED)]),
        (5.5, [C(f"사업자번호 {co.get('biz_no', '')} · 대표 {co.get('ceo', '')}", size=8.5, color=MUTED),
               C(doc_meta(doc), size=9, color=MUTED, align='RIGHT')]),
        (1, [C(cs=2, bg=ACCENT)]),
    ], border=False)]


def foot_text(doc, co, kind):
    d = doc['data']
    parts = []
    if d.get('note'):
        parts.append('비고\n' + d['note'])
    bank = doc_bank(doc, co)
    if bank:
        parts.append('입금계좌\n' + bank)
    if kind == 'statement':
        parts.append('인수자 확인   ________________ (인)')
    return '\n\n'.join(parts)


def letter_trade(doc, co, kind):
    d, c, K = doc['data'], doc['data'].get('client') or {}, TRADE[kind]
    items = trade_items(doc)
    supply, tax = sum_items(items, d.get('vat_mode'))
    meta = trade_meta(kind, doc, c, d) or [('', '')]
    to = [(10, [C(f"{doc['client_name']}   귀하", cs=2, size=14, bold=True), C(),
                C(f'합계 금액 ({vat_word(d)})', bg=SOFT, color=ACCENT, bold=True, size=9, align='RIGHT', valign='BOTTOM')])]
    for i, (k, v) in enumerate(meta):
        cells = [C(k, color=ACCENT, bold=True, size=9.5), C(v, size=9.5), C()]
        if i == 0:
            cells.append(C(f'₩ {won(supply + tax)}', rs=len(meta), bg=SOFT, color=ACCENT, bold=True, size=20, align='RIGHT'))
        to.append((6.5, cells))
    to.append((8, [C(K['greet'], cs=2, size=9, color=MUTED), C(), C(f'금 {kor_num(supply + tax)}원 정', bg=SOFT, size=8.5, align='RIGHT')]))
    widths = item_widths(K['numbered'])
    return letter_head(co, K['title'], doc) + [
        ('s', 5), T([22, 86, 6, 66], to, border=False),
        ('s', 5), T(widths, item_rows(items, d, K['numbered'], 'letter'), border=False),
        ('s', 1), T(widths, sum_rows(supply, tax, len(widths)), border=False),
        ('s', 6), T([100, 10, 54, 16], [
            (7, [C(foot_text(doc, co, kind), rs=3, size=9, valign='TOP'), C(rs=3), C(ymd(doc['doc_date']), cs=2, align='RIGHT')]),
            (8, [C(co['name'], cs=2, bold=True, size=12.5, align='RIGHT')]),
            (9, [C(f"대표  {co['ceo']}", align='RIGHT'), C('(인)', align='CENTER', stamp=True)]),
        ], border=False),
    ]


def card_top(co, plain, en, doc):
    return [T([120, 60], [
        (12, [C(plain, size=24, bold=True, color=ACCENT), C(img=('logo', 25, 7), align='RIGHT', valign='BOTTOM')]),
        (6, [C(doc_meta(doc, en), size=8.5, color=MUTED), C(co['name'], bold=True, size=12, align='RIGHT')]),
    ], border=False)]


def card_trade(doc, co, kind):
    d, c, K = doc['data'], doc['data'].get('client') or {}, TRADE[kind]
    items = trade_items(doc)
    supply, tax = sum_items(items, d.get('vat_mode'))
    total = supply + tax
    if kind == 'quote':
        sub1 = ' · '.join(x for x in [c.get('tel') and 'TEL ' + c['tel'], c.get('fax') and 'FAX ' + c['fax']] if x)
        sub2 = '유효기간 ' + d['valid_until'] if d.get('valid_until') else ''
    else:
        sub1 = ' · '.join(x for x in [c.get('biz_no') and '사업자 ' + c['biz_no'], c.get('ceo') and '대표 ' + c['ceo']] if x)
        sub1 += ('\n' + c['address']) if c.get('address') else ''
        sub2 = '연락처 ' + c['tel'] if c.get('tel') else ''
    soft = lambda t='', **kw: C(t, bg=SOFT, **kw)
    def acc(t='', b=(1, 1, 1, 1), **kw):  # 인쇄용: 채우지 않고 강조색 글씨 + 테두리
        return C(t, color=ACCENT, b=b, **kw)
    widths = item_widths(K['numbered'])
    return card_top(co, K['plain'], K['en'], doc) + [
        ('s', 5),
        T([57, 4, 57, 4, 58], [
            (6.5, [soft('받는 분' if kind == 'quote' else '공급받는자', color=ACCENT, bold=True, size=8.5), C(rs=3),
                   soft('공사명', color=ACCENT, bold=True, size=8.5), C(rs=3), acc(f'합계 금액 · {vat_word(d)}', b=(1, 1, 1, 0), bold=True, size=8.5)]),
            (9, [soft(doc['client_name'], bold=True, size=11), soft(doc.get('title'), bold=True, size=11), acc(f'₩{won(total)}', b=(1, 1, 0, 0), bold=True, size=17)]),
            (10, [soft(sub1, size=8.5, color=MUTED, valign='TOP'), soft(sub2, size=8.5, color=MUTED, valign='TOP'),
                  acc(f'금 {kor_num(total)}원 정', b=(1, 1, 0, 1), size=8.5, valign='TOP')]),
        ], border=False),
        ('s', 3), P(K['greet'], size=9.5, color=MUTED),
        ('s', 1), T(widths, item_rows(items, d, K['numbered'], 'card'), border=False),
        ('s', 3), T([79, 32, 2, 30, 2, 35], [(8, [C(), soft(f'공급가액  {won(supply)}', size=9, align='CENTER'), C(),
                                                   soft(f'부가세  {won(tax)}', size=9, align='CENTER'), C(),
                                                   acc(f'합계  ₩{won(total)}', bold=True, size=9.5, align='CENTER')])], border=False),
        ('s', 5),
        T([86, 4, 74, 16], [
            (6.5, [soft('비고 · 결제', color=ACCENT, bold=True, size=8.5), C(rs=4), soft('공급자', cs=2, color=ACCENT, bold=True, size=8.5)]),
            (8, [soft(foot_text(doc, co, kind), rs=3, size=9, valign='TOP'), soft(co['name'], cs=2, bold=True, size=11)]),
            (8, [soft(f"사업자번호 {co.get('biz_no', '')} · 대표 {co['ceo']}", size=8.5), soft('(인)', align='CENTER', stamp=True)]),
            (11, [soft(f"{co.get('address', '')}\nTEL {co.get('tel', '')}" + (f" · {co['email']}" if co.get('email') else ''),
                       cs=2, size=8.5, color=MUTED, valign='TOP')]),
        ], border=False),
    ]


def side_page(doc, co, main, ncols, widths):
    """사이드바형: 왼쪽 색 띠(회사 정보 · 직인) + 오른쪽 본문을 한 표로"""
    n = len(main) + 3
    main = main + [(9, [C(cs=ncols)]), (8, [C(cs=ncols)]), (10, [C(cs=ncols)])]  # 직인 자리
    bank = doc_bank(doc, co) if doc['data'].get('items') else ''
    info = [('대표', co.get('ceo')), ('사업자번호', co.get('biz_no')), ('주소', co.get('address')), ('전화', co.get('tel')),
            ('팩스', co.get('fax')), ('이메일', co.get('email')), ('입금계좌', bank)]
    info_text = '\n\n'.join(f'{k}\n{v}' for k, v in info if v)
    left = {0: [C(co['name'], cs=2, rs=2, bg=SOFT, color=ACCENT, bold=True, size=14, valign='TOP')],
            2: [C(info_text, cs=2, rs=n - 5, bg=SOFT, size=8.5, valign='TOP')],
            n - 3: [C(ymd(doc['doc_date']), cs=2, bg=WHITE, size=9, align='CENTER', b=(1, 1, 1, 0))],
            n - 2: [C(co['name'], cs=2, bg=WHITE, bold=True, size=10, align='CENTER', b=(1, 1, 0, 0))],
            n - 1: [C(f"대표 {co['ceo']}", bg=WHITE, size=9, align='RIGHT', b=(1, 0, 0, 1)),
                    C('(인)', bg=WHITE, size=9, align='CENTER', stamp=True, b=(0, 1, 0, 1))]}
    rows = []
    for i, (h, cells) in enumerate(main):
        rows.append((h, left.get(i, []) + ([C(rs=n)] if i == 0 else []) + cells))
    return T([34, 14, 4] + widths, rows, border=False)


def side_trade(doc, co, kind):
    d, c, K = doc['data'], doc['data'].get('client') or {}, TRADE[kind]
    items = trade_items(doc)
    supply, tax = sum_items(items, d.get('vat_mode'))
    widths = item_widths(K['numbered'], compact=True)
    n = len(widths)
    full = lambda t='', **kw: C(t, cs=n, **kw)
    main = [(12, [full(K['title'], size=21, bold=True, color=ACCENT)]),
            (6, [full(doc_meta(doc, K['en']), size=8.5, color=MUTED)]),
            (9, [full(f"{doc['client_name']}   귀하", size=13, bold=True, valign='BOTTOM')])]
    main += [(6, [full(f'{k}    {v}', size=9.5)]) for k, v in trade_meta(kind, doc, c, d)]
    main += [(3, [full()]),
             (6, [full(f'합계 금액 ({vat_word(d)})', bg=SOFT, color=ACCENT, bold=True, size=8.5)]),
             (10, [full(f'₩ {won(supply + tax)}', bg=SOFT, color=ACCENT, bold=True, size=17)]),
             (6, [full(f'금 {kor_num(supply + tax)}원 정', bg=SOFT, size=8.5)]),
             (8, [full(K['greet'], size=9, color=MUTED)])]
    main += item_rows(items, d, K['numbered'], 'side')
    main += sum_rows(supply, tax, n, size=9, grand=9.5)  # 좁은 칸이라 작게
    if d.get('note'):
        main.append((10, [full(f"비고    {d['note']}", size=9, valign='TOP')]))
    if kind == 'statement':
        main.append((9, [full('인수자 확인    ________________ (인)', size=9.5)]))
    return [side_page(doc, co, main, n, widths)]


def completion_info(doc):
    d = doc['data']
    amt = int(doc.get('supply_amount') or 0)
    return [('공사명', doc.get('title')), ('발주처', doc['client_name']),
            ('계약금액', f"일금 {kor_num(amt)}원정 (₩{won(amt)}, {vat_label(d.get('vat_mode'))})" if amt else ''),
            ('계약일자', ymd_short(d.get('contract_date'))), ('착공일자', ymd_short(d.get('start_date'))),
            ('준공일자', ymd_short(d.get('end_date')))]


def photo_count(d):
    if (d.get('photo_layout') or 'rows') == 'grid':
        return len(d.get('photos') or [])
    return sum(len(g.get('photos') or []) for g in d.get('groups') or [])


def completion_photos(doc, co):
    """작업사진 페이지 — 표형 작업완료서의 사진 부분을 그대로 사용"""
    blocks = layout_completion(doc, co)
    return blocks[next(i for i, b in enumerate(blocks) if b[0] == 'br'):] if any(b[0] == 'br' for b in blocks) else []


def completion_signs(doc, co, soft=False):
    box = (lambda b: {'bg': SOFT}) if soft else (lambda b: {'b': b})
    return T([86, 8, 70, 16], [
        (7, [C('발주처 확인', color=ACCENT, bold=True, size=9, **box((1, 1, 1, 0))), C(rs=3),
             C('시공사', cs=2, color=ACCENT, bold=True, size=9, **box((1, 1, 1, 0)))]),
        (9, [C(doc['client_name'], bold=True, size=12, **box((1, 1, 0, 0))), C(co['name'], cs=2, bold=True, size=12, **box((1, 1, 0, 0)))]),
        (11, [C('확인자   ______________ (인)', **box((1, 1, 0, 1))), C(f"대표  {co['ceo']}", align='RIGHT', **box((1, 0, 0, 1))),
              C('(인)', align='CENTER', stamp=True, **box((0, 1, 0, 1)))]),
    ], border=False)


def letter_completion(doc, co):
    d = doc['data']
    return letter_head(co, '작업완료확인서', doc) + [
        ('s', 9), P('아래 공사를 완료하였음을 확인합니다.', size=12),
        ('s', 3), T([32, 148], [(10, [C(k, color=ACCENT, bold=True, size=11, b=HAIR), C(v, size=11, b=HAIR)])
                                for k, v in completion_info(doc)], border=False),
        ('s', 7), P(f'별첨 : 작업사진 {photo_count(d)}매', size=10.5, color=MUTED),
        ('s', 12), P(ymd(doc['doc_date']), size=13, align='CENTER'),
        ('s', 10), completion_signs(doc, co),
    ] + completion_photos(doc, co)


def card_completion(doc, co):
    d = doc['data']
    info = completion_info(doc)
    rows = []
    for i in range(0, len(info), 2):
        (k1, v1), (k2, v2) = info[i], info[i + 1]
        first = i == 0
        rows += [(6, [C(k1, bg=SOFT, color=ACCENT, bold=True, size=8.5, **({'b': (1, 1, 1, 0)} if first else {})), C(),
                      C(k2, bg=SOFT, color=ACCENT, bold=True, size=8.5)]),
                 (11, [C(v1 or '-', bg=SOFT, color=ACCENT if first else None, bold=True, size=11, valign='TOP',
                         **({'b': (1, 1, 0, 1)} if first else {})), C(),
                       C(v2 or '-', bg=SOFT, bold=True, size=11, valign='TOP')]),
                 (4, [C(cs=3)])]
    return card_top(co, '작업완료확인서', 'COMPLETION', doc) + [
        ('s', 5), P('아래 공사를 완료하였음을 확인합니다.', size=10, color=MUTED),
        ('s', 3), T([88, 4, 88], rows, border=False),
        ('s', 5), P(f'별첨 : 작업사진 {photo_count(d)}매', size=10, color=MUTED),
        ('s', 10), P(ymd(doc['doc_date']), size=13, align='CENTER'),
        ('s', 8), completion_signs(doc, co, soft=True),
    ] + completion_photos(doc, co)


def side_completion(doc, co):
    d = doc['data']
    widths = [32, 96]
    main = [(12, [C('작업완료확인서', cs=2, size=21, bold=True, color=ACCENT)]),
            (6, [C(doc_meta(doc, 'COMPLETION'), cs=2, size=8.5, color=MUTED)]),
            (16, [C('아래 공사를 완료하였음을 확인합니다.', cs=2, size=12, valign='BOTTOM')])]
    main += [(10, [C(k, color=ACCENT, bold=True, size=11, b=HAIR), C(v, size=11, b=HAIR)]) for k, v in completion_info(doc)]
    main += [(14, [C(f'별첨 : 작업사진 {photo_count(d)}매', cs=2, size=10.5, color=MUTED)]),
             (8, [C('발주처 확인', cs=2, color=ACCENT, bold=True, size=9, b=(1, 1, 1, 0))]),
             (9, [C(doc['client_name'], cs=2, bold=True, size=12, b=(1, 1, 0, 0))]),
             (11, [C('확인자   ______________ (인)', cs=2, b=(1, 1, 0, 1))]),
             (60, [C(cs=2)])]
    return [side_page(doc, co, main, 2, widths)] + completion_photos(doc, co)


def pretty_contract(doc, co):
    """계약서: 표 형식 그대로, 구역 칸 · 항목 칸에 색을 넣어 깔끔하게"""
    blocks = layout_contract(doc, co)
    for _, cells in blocks[0][1]['rows']:
        for cell in cells:
            role = cell.get('role')
            if role == 'sec':
                cell.update(bg=SOFT, color=ACCENT, bold=True)
            elif role == 'label':
                cell.update(bg=LABBG, color=LABFG, bold=True)
            elif role == 'text':
                cell.update(bg=SOFT)
            elif role == 'who':
                cell.update(color=ACCENT, bold=True)
            if cell.get('title'):
                cell['title'] = False
                cell['color'] = ACCENT
    return blocks


NEW_LAYOUTS = {
    'letter': {'quote': lambda d, c: letter_trade(d, c, 'quote'), 'statement': lambda d, c: letter_trade(d, c, 'statement'),
               'contract': pretty_contract, 'completion': letter_completion},
    'card': {'quote': lambda d, c: card_trade(d, c, 'quote'), 'statement': lambda d, c: card_trade(d, c, 'statement'),
             'contract': pretty_contract, 'completion': card_completion},
    'sidebar': {'quote': lambda d, c: side_trade(d, c, 'quote'), 'statement': lambda d, c: side_trade(d, c, 'statement'),
                'contract': pretty_contract, 'completion': side_completion},
}


# 문서 디자인 시안 — print.html의 .theme-* 와 같은 색 · 글꼴
THEMES = {
    'classic': {'lab_bg': '#EFEFEF', 'lab_fg': '#000000', 'line': '#777777', 'accent': '#000000', 'soft': '#DFE5EE',
                'underline': True, 'title_align': 'CENTER', 'font_xlsx': '맑은 고딕', 'font_hwpx': '함초롬돋움'},
    'blue':    {'lab_bg': '#E3EDF9', 'lab_fg': '#0B3F7A', 'line': '#8FA9C9', 'accent': '#0B4F9C', 'soft': '#E3EDF9',
                'underline': False, 'title_align': 'CENTER', 'font_xlsx': '맑은 고딕', 'font_hwpx': '함초롬돋움'},
    'modern':  {'lab_bg': '#E9ECEF', 'lab_fg': '#2F3640', 'line': '#9AA0A6', 'accent': '#2F3640', 'soft': '#F1F3F5',
                'underline': False, 'title_align': 'LEFT', 'font_xlsx': '맑은 고딕', 'font_hwpx': '함초롬돋움'},
    'serif':   {'lab_bg': '#E8F1EC', 'lab_fg': '#1F4D3A', 'line': '#8AA898', 'accent': '#1F4D3A', 'soft': '#E8F1EC',
                'underline': False, 'title_align': 'CENTER', 'font_xlsx': '바탕', 'font_hwpx': '함초롬바탕'},
}


def theme_of(company):
    return THEMES.get((company or {}).get('theme'), THEMES['classic'])


def apply_theme(blocks, th):
    """레이아웃의 기본 색(회색 머리칸 · 파란 바탕 · 제목 · @자리표시)을 시안 색으로 바꾼다"""
    slot = {ACCENT: th['accent'], SOFT: th['soft'], LABBG: th['lab_bg'], LABFG: th['lab_fg']}
    for kind, *rest in blocks:
        if kind == 'p':
            p = rest[0]
            if p.get('color') in slot:
                p['color'] = slot[p['color']]
            if p.get('title'):
                p.update(color=th['accent'], underline=th['underline'], align=th['title_align'])
            elif p.get('text') == '작업사진':
                p['color'] = th['accent']
        elif kind == 't':
            for _, cells in rest[0]['rows']:
                for cell in cells:
                    for k in ('bg', 'color'):
                        if cell.get(k) in slot:
                            cell[k] = slot[cell[k]]
                    if cell.get('bg') == GRAY:
                        cell.update(bg=th['lab_bg'], color=th['lab_fg'])
                    elif cell.get('bg') == BLUE:
                        cell['bg'] = th['soft']
                    if cell.get('title'):
                        cell['color'] = th['accent']
    return blocks


LAYOUTS = {'quote': layout_quote, 'statement': layout_statement, 'contract': layout_contract, 'completion': layout_completion}


def build_blocks(doc, co):
    table = NEW_LAYOUTS.get(co.get('layout'), LAYOUTS)  # 표형(기본)은 LAYOUTS
    fn = table[doc['doc_type']]
    return apply_theme(fn(doc, co), theme_of(co))


# ── 이미지 ───────────────────────────────────────────────────────────────────

class Images:
    """레이아웃의 img 키 → 실제 이미지 바이트 (사진은 칸 비율에 맞게 잘라서 축소)"""

    def __init__(self, photo_dir, stamp_path, logo_path):
        self.photo_dir, self.stamp_path, self.logo_path = photo_dir, stamp_path, logo_path
        self.cache = {}

    def get(self, key):
        """반환: (bytes, ext, px_w, px_h) 또는 None"""
        if key in self.cache:
            return self.cache[key]
        res = None
        try:
            if key[0] == 'stamp':
                res = self._png(self.stamp_path)
            elif key[0] == 'logo':
                res = self._png(self.logo_path)
            elif key[0] == 'photo':
                _, folder, filename, w, h = key
                path = os.path.join(self.photo_dir, os.path.basename(folder), os.path.basename(filename))
                img = ImageOps.exif_transpose(Image.open(path)).convert('RGB')
                img = ImageOps.fit(img, (int(w * 12), int(h * 12)), Image.LANCZOS)  # 약 300dpi
                buf = io.BytesIO()
                img.save(buf, 'JPEG', quality=82)
                res = (buf.getvalue(), 'jpg', img.width, img.height)
        except Exception:
            res = None
        self.cache[key] = res
        return res

    @staticmethod
    def _png(path):
        if not path or not os.path.exists(path):
            return None
        img = Image.open(path)
        buf = io.BytesIO()
        img.save(buf, 'PNG')
        return (buf.getvalue(), 'png', img.width, img.height)


def place_cells(rows):
    """병합 칸을 고려해 각 셀의 (행, 열) 위치 계산"""
    occupied, placed = set(), []
    for r, (_, cells) in enumerate(rows):
        col = 0
        for cell in cells:
            while (r, col) in occupied:
                col += 1
            cs, rs = cell.get('cs', 1), cell.get('rs', 1)
            for dr in range(rs):
                for dc in range(cs):
                    occupied.add((r + dr, col + dc))
            placed.append((r, col, cell))
            col += cs
    return placed


def cell_borders(cell, table_border):
    b = cell.get('b')
    if b is None:
        b = (1, 1, 1, 1) if table_border else (0, 0, 0, 0)
    return tuple(bool(x) for x in b)


def cell_text(cell):
    if 'num' in cell:
        return won(cell['num'])
    return cell.get('text', '')


# ── 엑셀 ─────────────────────────────────────────────────────────────────────

def render_xlsx(blocks, images, title, theme=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.utils import get_column_letter

    PX = 96 / 25.4   # px per mm
    EMU = 9525       # EMU per px
    th = theme or THEMES['classic']
    FONT = th['font_xlsx']
    hexc = lambda c: (c or '#000000').lstrip('#')

    # 모든 표의 열 경계를 합쳐 시트 전체의 열 격자를 만든다
    edges = {0.0, float(PAGE_W)}
    for kind, *rest in blocks:
        if kind == 't':
            x = 0.0
            for w in rest[0]['widths']:
                x = round(x + w, 2)
                edges.add(x)
    edges = sorted(edges)
    col_of = {x: i for i, x in enumerate(edges)}  # 경계 x(mm) → 열 번호(0부터)

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.sheet_view.showGridLines = False
    for i in range(len(edges) - 1):
        px = (edges[i + 1] - edges[i]) * PX
        ws.column_dimensions[get_column_letter(i + 1)].width = max(0.1, (px - 5) / 7)

    thin = Side(style='thin', color=hexc(th['line']))
    row = 1
    last_col = len(edges) - 1

    def add_image(key, w_mm, h_mm, x_mm, row_idx, y_mm):
        im = images.get(key)
        if not im:
            return
        x_mm = max(0.0, min(x_mm, PAGE_W - 0.5))
        c = max(i for i, e in enumerate(edges[:-1]) if e <= x_mm)
        # 칸 위쪽으로 넘치면 윗행 기준으로 위치를 잡는다
        while y_mm < 0 and row_idx > 1:
            row_idx -= 1
            y_mm += (ws.row_dimensions[row_idx].height or 15) * 25.4 / 72
        xl = XLImage(io.BytesIO(im[0]))
        size = XDRPositiveSize2D(int(w_mm * PX * EMU), int(h_mm * PX * EMU))
        xl.anchor = OneCellAnchor(_from=AnchorMarker(col=c, colOff=int((x_mm - edges[c]) * PX * EMU),
                                                     row=row_idx - 1, rowOff=int(max(0, y_mm) * PX * EMU)), ext=size)
        ws.add_image(xl)

    for kind, *rest in blocks:
        if kind == 's':
            ws.row_dimensions[row].height = rest[0] * 72 / 25.4
            row += 1
        elif kind == 'br':
            ws.row_breaks.append(Break(id=row - 1))
        elif kind == 'p':
            p = rest[0]
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
            cell = ws.cell(row=row, column=1, value=p['text'])
            cell.font = Font(name=FONT, size=p.get('size', 10), bold=p.get('bold', False), color=hexc(p.get('color')),
                             underline='single' if p.get('underline') else None)
            cell.alignment = Alignment(horizontal=p.get('align', 'LEFT').lower(), vertical='center')
            ws.row_dimensions[row].height = p.get('size', 10) * 1.8
            row += 1
        elif kind == 't':
            t = rest[0]
            xs = [0.0]
            for w in t['widths']:
                xs.append(round(xs[-1] + w, 2))
            heights = [h for h, _ in t['rows']]
            for i, h in enumerate(heights):
                ws.row_dimensions[row + i].height = h * 72 / 25.4
            for r, c, cell in place_cells(t['rows']):
                cs, rs = cell.get('cs', 1), cell.get('rs', 1)
                c0, c1 = col_of[xs[c]] + 1, col_of[xs[c + cs]]
                r0, r1 = row + r, row + r + rs - 1
                if c1 > c0 or r1 > r0:
                    ws.merge_cells(start_row=r0, start_column=c0, end_row=r1, end_column=c1)
                xc = ws.cell(row=r0, column=c0)
                if 'num' in cell:
                    xc.value = int(cell['num'])
                    xc.number_format = '#,##0'
                elif cell.get('text'):
                    v = cell['text']
                    xc.value = int(v) if isinstance(v, str) and v.isdigit() and len(v) < 6 else v
                xc.font = Font(name=FONT, size=cell.get('size', 10), bold=cell.get('bold', False), color=hexc(cell.get('color')))
                xc.alignment = Alignment(horizontal=cell.get('align', 'LEFT').lower(),
                                         vertical={'TOP': 'top', 'BOTTOM': 'bottom'}.get(cell.get('valign'), 'center'),
                                         wrap_text=True, indent=0 if cell.get('align') in ('CENTER', 'RIGHT') else 1)
                fill = PatternFill('solid', fgColor=cell['bg'][1:]) if cell.get('bg') else None
                bl, br_, bt, bb = cell_borders(cell, t['border'])
                for rr in range(r0, r1 + 1):
                    for cc in range(c0, c1 + 1):
                        x = ws.cell(row=rr, column=cc)
                        x.border = Border(left=thin if bl and cc == c0 else None, right=thin if br_ and cc == c1 else None,
                                          top=thin if bt and rr == r0 else None, bottom=thin if bb and rr == r1 else None)
                        if fill:
                            x.fill = fill
                # 셀 안 이미지 (로고 · 사진)
                cell_w = xs[c + cs] - xs[c]
                cell_h = sum(heights[r:r + rs])
                if cell.get('img'):
                    key = cell['img']
                    w, h = key[-2], key[-1]
                    if cell.get('align') == 'RIGHT':
                        x0 = xs[c] + cell_w - w - 1
                    else:
                        x0 = xs[c] + (cell_w - w) / 2
                    y0 = {'BOTTOM': cell_h - h - 0.5, 'TOP': 0.5}.get(cell.get('valign'), (cell_h - h) / 2)
                    add_image(key, w, h, x0, r0, y0)
                # 직인: (인) 칸 정중앙에 겹쳐 찍기
                if cell.get('stamp'):
                    add_image(('stamp',), STAMP, STAMP, xs[c] + cell_w / 2 - STAMP / 2, r0, cell_h / 2 - STAMP / 2)
            row += len(t['rows'])

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0 if ws.row_breaks.brk else 1
    ws.page_margins.left = ws.page_margins.right = 0.55
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.4
    ws.page_margins.header = ws.page_margins.footer = 0
    ws.print_options.horizontalCentered = True
    ws.print_area = f'A1:{get_column_letter(last_col)}{row - 1}'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 한글(.hwpx) ──────────────────────────────────────────────────────────────

HU = 7200 / 25.4  # HWPUNIT per mm
NS = ('xmlns:ha="http://www.hancom.co.kr/hwpml/2011/app" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" '
      'xmlns:hp10="http://www.hancom.co.kr/hwpml/2016/paragraph" xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
      'xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core" xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head" '
      'xmlns:hhs="http://www.hancom.co.kr/hwpml/2011/history" xmlns:hm="http://www.hancom.co.kr/hwpml/2011/master-page" '
      'xmlns:hpf="http://www.hancom.co.kr/schema/2011/hpf" xmlns:dc="http://purl.org/dc/elements/1.1/" '
      'xmlns:opf="http://www.idpf.org/2007/opf/" xmlns:ooxmlchart="http://www.hancom.co.kr/hwpml/2016/ooxmlchart" '
      'xmlns:hwpunitchar="http://www.hancom.co.kr/hwpml/2016/HwpUnitChar" xmlns:epub="http://www.idpf.org/2007/ops" '
      'xmlns:config="urn:oasis:names:tc:opendocument:xmlns:config:1.0"')
LANGS = ['HANGUL', 'LATIN', 'HANJA', 'JAPANESE', 'OTHER', 'SYMBOL', 'USER']
XML_HEAD = '<?xml version="1.0" encoding="UTF-8" standalone="yes" ?>'


def hu(mm):
    return int(round(mm * HU))


class Hwpx:
    def __init__(self, images, page_border=False, theme=None):
        self.images = images
        self.theme = theme or THEMES['classic']
        self.char_prs = {}    # (size, bold, underline, color) → id
        self.para_prs = {}    # (align, line%) → id
        self.border_fills = {}  # (l, r, t, b, bg, line) → id  (1·2번은 기본값)
        self.bin_items = {}   # img key → (item id, bytes, ext, px_w, px_h)
        self.uid = random.randint(1000000, 9000000)
        self.page_border = page_border
        self.cpr(10)
        self.ppr('JUSTIFY', 160)

    def next_id(self):
        self.uid += 1
        return self.uid

    def cpr(self, size=10, bold=False, underline=False, color='#000000'):
        key = (float(size), bool(bold), bool(underline), color or '#000000')
        return self.char_prs.setdefault(key, len(self.char_prs))

    def ppr(self, align='LEFT', line=130):
        key = (align, int(line))
        return self.para_prs.setdefault(key, len(self.para_prs))

    def bfill(self, l, r, t, b, bg=None, line='SOLID', width='0.12 mm'):
        key = (l, r, t, b, bg, line, width)
        return self.border_fills.setdefault(key, len(self.border_fills) + 3)

    def bin_id(self, key):
        if key not in self.bin_items:
            im = self.images.get(key)
            if not im:
                return None
            self.bin_items[key] = (f'image{len(self.bin_items) + 1}',) + im
        return self.bin_items[key][0]

    # ── 본문 요소 ──
    def pic(self, key, w_mm, h_mm, inline=True):
        ref = self.bin_id(key)
        if not ref:
            return ''
        _, _, _, pw, ph = self.bin_items[key]
        W, H = hu(w_mm), hu(h_mm)
        pos = (f'<hp:pos treatAsChar="1" affectLSpacing="0" flowWithText="1" allowOverlap="0" holdAnchorAndSO="0" '
               f'vertRelTo="PARA" horzRelTo="COLUMN" vertAlign="TOP" horzAlign="LEFT" vertOffset="0" horzOffset="0"/>') if inline else (
               # 글 앞으로 띄워서 문단(=칸) 가운데에 놓는다
               '<hp:pos treatAsChar="0" affectLSpacing="0" flowWithText="0" allowOverlap="1" holdAnchorAndSO="0" '
               'vertRelTo="PARA" horzRelTo="PARA" vertAlign="CENTER" horzAlign="CENTER" vertOffset="0" horzOffset="0"/>')
        wrap = 'TOP_AND_BOTTOM' if inline else 'IN_FRONT_OF_TEXT'
        return (f'<hp:pic id="{self.next_id()}" zOrder="{self.next_id() % 1000}" numberingType="PICTURE" textWrap="{wrap}" '
                f'textFlow="BOTH_SIDES" lock="0" dropcapstyle="None" href="" groupLevel="0" instid="{self.next_id()}" reverse="0">'
                f'<hp:offset x="0" y="0"/><hp:orgSz width="{W}" height="{H}"/><hp:curSz width="{W}" height="{H}"/>'
                f'<hp:flip horizontal="0" vertical="0"/><hp:rotationInfo angle="0" centerX="{W // 2}" centerY="{H // 2}" rotateimage="1"/>'
                f'<hp:renderingInfo><hc:transMatrix e1="1" e2="0" e3="0" e4="0" e5="1" e6="0"/>'
                f'<hc:scaMatrix e1="1" e2="0" e3="0" e4="0" e5="1" e6="0"/><hc:rotMatrix e1="1" e2="0" e3="0" e4="0" e5="1" e6="0"/></hp:renderingInfo>'
                f'<hc:img binaryItemIDRef="{ref}" bright="0" contrast="0" effect="REAL_PIC" alpha="0"/>'
                f'<hp:imgRect><hc:pt0 x="0" y="0"/><hc:pt1 x="{W}" y="0"/><hc:pt2 x="{W}" y="{H}"/><hc:pt3 x="0" y="{H}"/></hp:imgRect>'
                f'<hp:imgClip left="0" right="{pw * 75}" top="0" bottom="{ph * 75}"/><hp:inMargin left="0" right="0" top="0" bottom="0"/>'
                f'<hp:imgDim dimwidth="{pw * 75}" dimheight="{ph * 75}"/><hp:effects/>'
                f'<hp:sz width="{W}" widthRelTo="ABSOLUTE" height="{H}" heightRelTo="ABSOLUTE" protect="0"/>{pos}'
                f'<hp:outMargin left="0" right="0" top="0" bottom="0"/></hp:pic>')

    def para(self, text, size=10, bold=False, underline=False, align='LEFT', line=130, extra='', page_break=False, color='#000000'):
        run = f'<hp:t>{escape(text)}</hp:t>' if text else ''
        return (f'<hp:p id="2147483648" paraPrIDRef="{self.ppr(align, line)}" styleIDRef="0" pageBreak="{1 if page_break else 0}" '
                f'columnBreak="0" merged="0"><hp:run charPrIDRef="{self.cpr(size, bold, underline, color)}">{run}{extra}</hp:run></hp:p>')

    def table(self, t, page_break=False):
        widths, rows = t['widths'], t['rows']
        xs = [0.0]
        for w in widths:
            xs.append(xs[-1] + w)
        heights = [h for h, _ in rows]
        tcs = {r: [] for r in range(len(rows))}
        for r, c, cell in place_cells(rows):
            cs, rs = cell.get('cs', 1), cell.get('rs', 1)
            cw, ch = xs[c + cs] - xs[c], sum(heights[r:r + rs])
            l, rr, tt, bb = cell_borders(cell, t['border'])
            bf = self.bfill(l, rr, tt, bb, cell.get('bg'))
            size, bold, align = cell.get('size', 10), cell.get('bold', False), cell.get('align', 'LEFT')
            line = int(cell.get('line', 1.3) * 100)
            lines = cell_text(cell).split('\n')
            paras = []
            for i, ln in enumerate(lines):
                extra = self.pic(('stamp',), STAMP, STAMP, inline=False) if cell.get('stamp') and i == 0 else ''
                paras.append(self.para(ln, size, bold, False, align, line, extra, color=cell.get('color')))
            if cell.get('img'):
                key = cell['img']
                paras = [self.para('', size, False, False, align, 100, self.pic(key, key[-2], key[-1]))]
            valign = {'TOP': 'TOP', 'BOTTOM': 'BOTTOM'}.get(cell.get('valign'), 'CENTER')
            pad_lr = 0.4 if cell.get('img') else 1.5
            tcs[r].append(
                f'<hp:tc name="" header="0" hasMargin="1" protect="0" editable="0" dirty="0" borderFillIDRef="{bf}">'
                f'<hp:subList id="" textDirection="HORIZONTAL" lineWrap="BREAK" vertAlign="{valign}" linkListIDRef="0" '
                f'linkListNextIDRef="0" textWidth="0" textHeight="0" hasTextRef="0" hasNumRef="0">{"".join(paras)}</hp:subList>'
                f'<hp:cellAddr colAddr="{c}" rowAddr="{r}"/><hp:cellSpan colSpan="{cs}" rowSpan="{rs}"/>'
                f'<hp:cellSz width="{hu(cw)}" height="{hu(ch)}"/>'
                f'<hp:cellMargin left="{hu(pad_lr)}" right="{hu(pad_lr)}" top="{hu(0.4)}" bottom="{hu(0.4)}"/></hp:tc>')
        trs = ''.join(f'<hp:tr>{"".join(tcs[r])}</hp:tr>' for r in range(len(rows)))
        tbl = (f'<hp:tbl id="{self.next_id()}" zOrder="{self.next_id() % 1000}" numberingType="TABLE" textWrap="TOP_AND_BOTTOM" '
               f'textFlow="BOTH_SIDES" lock="0" dropcapstyle="None" pageBreak="CELL" repeatHeader="0" rowCnt="{len(rows)}" '
               f'colCnt="{len(widths)}" cellSpacing="0" borderFillIDRef="{self.bfill(0, 0, 0, 0)}" noAdjust="0">'
               f'<hp:sz width="{hu(xs[-1])}" widthRelTo="ABSOLUTE" height="{hu(sum(heights))}" heightRelTo="ABSOLUTE" protect="0"/>'
               f'<hp:pos treatAsChar="1" affectLSpacing="0" flowWithText="1" allowOverlap="0" holdAnchorAndSO="0" vertRelTo="PARA" '
               f'horzRelTo="COLUMN" vertAlign="TOP" horzAlign="LEFT" vertOffset="0" horzOffset="0"/>'
               f'<hp:outMargin left="0" right="0" top="0" bottom="0"/>'
               f'<hp:inMargin left="{hu(1.5)}" right="{hu(1.5)}" top="{hu(0.4)}" bottom="{hu(0.4)}"/>{trs}</hp:tbl>')
        return (f'<hp:p id="2147483648" paraPrIDRef="{self.ppr("LEFT", 100)}" styleIDRef="0" pageBreak="{1 if page_break else 0}" '
                f'columnBreak="0" merged="0"><hp:run charPrIDRef="{self.cpr(1)}">{tbl}<hp:t/></hp:run></hp:p>')

    def body(self, blocks):
        out, brk = [], False
        for kind, *rest in blocks:
            if kind == 'br':
                brk = True
                continue
            if kind == 's':
                # 빈 줄 하나로 여백 (줄 높이 ≈ 글자 크기)
                out.append(self.para('', max(1, rest[0] / 0.3528), line=100, page_break=brk))
            elif kind == 'p':
                p = rest[0]
                out.append(self.para(p['text'], p.get('size', 10), p.get('bold', False), p.get('underline', False),
                                     p.get('align', 'LEFT'), 130, page_break=brk, color=p.get('color')))
            elif kind == 't':
                out.append(self.table(rest[0], page_break=brk))
            brk = False
        return out

    # ── 패키지 ──
    def sec_pr(self):
        pbf = self.bfill(1, 1, 1, 1, None, 'DOUBLE_SLIM', '0.7 mm') if self.page_border else 1
        pbf_xml = ''.join(
            f'<hp:pageBorderFill type="{t}" borderFillIDRef="{pbf}" textBorder="PAPER" headerInside="0" footerInside="0" fillArea="PAPER">'
            f'<hp:offset left="{hu(7)}" right="{hu(7)}" top="{hu(7)}" bottom="{hu(7)}"/></hp:pageBorderFill>' for t in ('BOTH', 'EVEN', 'ODD'))
        return ('<hp:secPr id="" textDirection="HORIZONTAL" spaceColumns="1134" tabStop="8000" tabStopVal="4000" tabStopUnit="HWPUNIT" '
                'outlineShapeIDRef="1" memoShapeIDRef="0" textVerticalWidthHead="0" masterPageCnt="0">'
                '<hp:grid lineGrid="0" charGrid="0" wonggojiFormat="0"/><hp:startNum pageStartsOn="BOTH" page="0" pic="0" tbl="0" equation="0"/>'
                '<hp:visibility hideFirstHeader="0" hideFirstFooter="0" hideFirstMasterPage="0" border="SHOW_ALL" fill="SHOW_ALL" '
                'hideFirstPageNum="0" hideFirstEmptyLine="0" showLineNumber="0"/>'
                '<hp:lineNumberShape restartType="0" countBy="0" distance="0" startNumber="0"/>'
                f'<hp:pagePr landscape="WIDELY" width="59528" height="84186" gutterType="LEFT_ONLY">'
                f'<hp:margin header="0" footer="0" gutter="0" left="{hu(15)}" right="{hu(15)}" top="{hu(14)}" bottom="{hu(10)}"/></hp:pagePr>'
                '<hp:footNotePr><hp:autoNumFormat type="DIGIT" userChar="" prefixChar="" suffixChar=")" supscript="0"/>'
                '<hp:noteLine length="-1" type="SOLID" width="0.12 mm" color="#000000"/><hp:noteSpacing betweenNotes="283" belowLine="567" aboveLine="850"/>'
                '<hp:numbering type="CONTINUOUS" newNum="1"/><hp:placement place="EACH_COLUMN" beneathText="0"/></hp:footNotePr>'
                '<hp:endNotePr><hp:autoNumFormat type="DIGIT" userChar="" prefixChar="" suffixChar=")" supscript="0"/>'
                '<hp:noteLine length="14692344" type="SOLID" width="0.12 mm" color="#000000"/><hp:noteSpacing betweenNotes="0" belowLine="567" aboveLine="850"/>'
                '<hp:numbering type="CONTINUOUS" newNum="1"/><hp:placement place="END_OF_DOCUMENT" beneathText="0"/></hp:endNotePr>'
                f'{pbf_xml}</hp:secPr>')

    def header_xml(self):
        font = (f'<hh:font id="0" face="{self.theme["font_hwpx"]}" type="TTF" isEmbedded="0"><hh:typeInfo familyType="FCAT_GOTHIC" weight="6" '
                'proportion="4" contrast="0" strokeVariation="1" armStyle="1" letterform="1" midline="1" xHeight="1"/></hh:font>')
        fontfaces = ''.join(f'<hh:fontface lang="{lg}" fontCnt="1">{font}</hh:fontface>' for lg in LANGS)

        def bf_xml(i, l, r, t, b, bg, line='SOLID', width='0.12 mm'):
            # 표 선은 시안의 선 색, 쪽 테두리(이중선)는 시안의 강조색
            border_color = self.theme['accent'] if line == 'DOUBLE_SLIM' else self.theme['line']
            side = lambda on: f'type="{line if on else "NONE"}" width="{width if on else "0.1 mm"}" color="{border_color}"'
            fill = (f'<hc:fillBrush><hc:winBrush faceColor="{bg}" hatchColor="#999999" alpha="0"/></hc:fillBrush>' if bg else '')
            return (f'<hh:borderFill id="{i}" threeD="0" shadow="0" centerLine="NONE" breakCellSeparateLine="0">'
                    '<hh:slash type="NONE" Crooked="0" isCounter="0"/><hh:backSlash type="NONE" Crooked="0" isCounter="0"/>'
                    f'<hh:leftBorder {side(l)}/><hh:rightBorder {side(r)}/><hh:topBorder {side(t)}/><hh:bottomBorder {side(b)}/>'
                    f'<hh:diagonal type="SOLID" width="0.1 mm" color="#000000"/>{fill}</hh:borderFill>')
        fills = [bf_xml(1, 0, 0, 0, 0, None), bf_xml(2, 0, 0, 0, 0, None)]
        fills += [bf_xml(i, *k) for k, i in sorted(self.border_fills.items(), key=lambda x: x[1])]

        attrs7 = lambda v: ' '.join(f'{lg.lower()}="{v}"' for lg in LANGS)
        chars = []
        for (size, bold, underline, color), i in sorted(self.char_prs.items(), key=lambda x: x[1]):
            chars.append(f'<hh:charPr id="{i}" height="{int(size * 100)}" textColor="{color}" shadeColor="none" useFontSpace="0" '
                         f'useKerning="0" symMark="NONE" borderFillIDRef="2"><hh:fontRef {attrs7(0)}/><hh:ratio {attrs7(100)}/>'
                         f'<hh:spacing {attrs7(0)}/><hh:relSz {attrs7(100)}/><hh:offset {attrs7(0)}/>{"<hh:bold/>" if bold else ""}'
                         f'<hh:underline type="{"BOTTOM" if underline else "NONE"}" shape="SOLID" color="{color}"/>'
                         '<hh:strikeout shape="NONE" color="#000000"/><hh:outline type="NONE"/>'
                         '<hh:shadow type="NONE" color="#C0C0C0" offsetX="10" offsetY="10"/></hh:charPr>')
        margin = ('<hh:margin><hc:intent value="0" unit="HWPUNIT"/><hc:left value="0" unit="HWPUNIT"/><hc:right value="0" unit="HWPUNIT"/>'
                  '<hc:prev value="0" unit="HWPUNIT"/><hc:next value="0" unit="HWPUNIT"/></hh:margin>')
        paras = []
        for (align, line), i in sorted(self.para_prs.items(), key=lambda x: x[1]):
            sp = f'{margin}<hh:lineSpacing type="PERCENT" value="{line}" unit="HWPUNIT"/>'
            paras.append(f'<hh:paraPr id="{i}" tabPrIDRef="0" condense="0" fontLineHeight="0" snapToGrid="0" suppressLineNumbers="0" checked="0">'
                         f'<hh:align horizontal="{align}" vertical="BASELINE"/><hh:heading type="NONE" idRef="0" level="0"/>'
                         '<hh:breakSetting breakLatinWord="KEEP_WORD" breakNonLatinWord="KEEP_WORD" widowOrphan="0" keepWithNext="0" '
                         'keepLines="0" pageBreakBefore="0" lineWrap="BREAK"/><hh:autoSpacing eAsianEng="0" eAsianNum="0"/>'
                         f'<hp:switch><hp:case hp:required-namespace="http://www.hancom.co.kr/hwpml/2016/HwpUnitChar">{sp}</hp:case>'
                         f'<hp:default>{sp}</hp:default></hp:switch>'
                         '<hh:border borderFillIDRef="2" offsetLeft="0" offsetRight="0" offsetTop="0" offsetBottom="0" connect="0" ignoreMargin="0"/>'
                         '</hh:paraPr>')
        return (f'{XML_HEAD}<hh:head {NS} version="1.4" secCnt="1">'
                '<hh:beginNum page="1" footnote="1" endnote="1" pic="1" tbl="1" equation="1"/><hh:refList>'
                f'<hh:fontfaces itemCnt="{len(LANGS)}">{fontfaces}</hh:fontfaces>'
                f'<hh:borderFills itemCnt="{len(fills)}">{"".join(fills)}</hh:borderFills>'
                f'<hh:charProperties itemCnt="{len(chars)}">{"".join(chars)}</hh:charProperties>'
                '<hh:tabProperties itemCnt="1"><hh:tabPr id="0" autoTabLeft="0" autoTabRight="0"/></hh:tabProperties>'
                f'<hh:paraProperties itemCnt="{len(paras)}">{"".join(paras)}</hh:paraProperties>'
                '<hh:styles itemCnt="1"><hh:style id="0" type="PARA" name="바탕글" engName="Normal" paraPrIDRef="0" charPrIDRef="0" '
                'nextStyleIDRef="0" langID="1042" lockForm="0"/></hh:styles>'
                '</hh:refList><hh:compatibleDocument targetProgram="HWP201X"><hh:layoutCompatibility/></hh:compatibleDocument>'
                '<hh:docOption><hh:linkinfo path="" pageInherit="0" footnoteInherit="0"/></hh:docOption>'
                '<hh:trackchageConfig flags="56"/></hh:head>')

    def build(self, blocks, title):
        paras = self.body(blocks)
        # 첫 문단에 구역 설정(용지 · 여백 · 쪽 테두리)을 넣는다
        first = (f'<hp:p id="2147483648" paraPrIDRef="0" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0">'
                 f'<hp:run charPrIDRef="{self.cpr(1)}">{self.sec_pr()}<hp:ctrl><hp:colPr id="" type="NEWSPAPER" layout="LEFT" '
                 'colCount="1" sameSz="1" sameGap="0"/></hp:ctrl></hp:run></hp:p>')
        section = f'{XML_HEAD}<hs:sec {NS}>{first}{"".join(paras)}</hs:sec>'
        header = self.header_xml()  # 본문을 만든 뒤에 (스타일이 모두 등록된 상태)
        bins = list(self.bin_items.values())
        manifest = ''.join(f'<opf:item id="{b[0]}" href="BinData/{b[0]}.{b[2]}" media-type="image/{"jpeg" if b[2] == "jpg" else "png"}" '
                           'isEmbeded="1"/>' for b in bins)
        hpf = (f'{XML_HEAD}<opf:package {NS} version="" unique-identifier="" id=""><opf:metadata><opf:title>{escape(title)}</opf:title>'
               '<opf:language>ko</opf:language></opf:metadata><opf:manifest>'
               '<opf:item id="header" href="Contents/header.xml" media-type="application/xml"/>'
               '<opf:item id="section0" href="Contents/section0.xml" media-type="application/xml"/>'
               f'<opf:item id="settings" href="settings.xml" media-type="application/xml"/>{manifest}</opf:manifest>'
               '<opf:spine><opf:itemref idref="header" linear="yes"/><opf:itemref idref="section0" linear="yes"/></opf:spine></opf:package>')
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr(zipfile.ZipInfo('mimetype'), 'application/hwp+zip', compress_type=zipfile.ZIP_STORED)
            z.writestr('version.xml', f'{XML_HEAD}<hv:HCFVersion xmlns:hv="http://www.hancom.co.kr/hwpml/2011/version" '
                       'tagetApplication="WORDPROCESSOR" major="5" minor="0" micro="5" buildNumber="0" os="1" xmlVersion="1.4" '
                       'application="Hancom Office Hangul" appVersion="9, 1, 1, 5656 WIN32LEWindows_8"/>')
            z.writestr('META-INF/container.xml', f'{XML_HEAD}<ocf:container xmlns:ocf="urn:oasis:names:tc:opendocument:xmlns:container" '
                       'xmlns:hpf="http://www.hancom.co.kr/schema/2011/hpf"><ocf:rootfiles><ocf:rootfile full-path="Contents/content.hpf" '
                       'media-type="application/hwpml-package+xml"/></ocf:rootfiles></ocf:container>')
            z.writestr('META-INF/manifest.xml', f'{XML_HEAD}<odf:manifest xmlns:odf="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"/>')
            z.writestr('Contents/content.hpf', hpf)
            z.writestr('Contents/header.xml', header)
            z.writestr('Contents/section0.xml', section)
            z.writestr('settings.xml', f'{XML_HEAD}<ha:HWPApplicationSetting xmlns:ha="http://www.hancom.co.kr/hwpml/2011/app" '
                       'xmlns:config="urn:oasis:names:tc:opendocument:xmlns:config:1.0"><ha:CaretPosition listIDRef="0" paraIDRef="0" pos="0"/>'
                       '</ha:HWPApplicationSetting>')
            for item_id, data, ext, _, _ in bins:
                z.writestr(f'BinData/{item_id}.{ext}', data)
        return buf.getvalue()


def render_hwpx(blocks, images, title, page_border=False, theme=None):
    return Hwpx(images, page_border, theme).build(blocks, title)
